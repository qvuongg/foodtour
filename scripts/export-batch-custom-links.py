#!/usr/bin/env python3
"""
Công cụ xuất danh sách link ShopeeFood từ crawler_urls_pending_affiliate.csv
sang các file Excel 'Batch Custom Links' chuẩn mẫu Shopee Affiliate.

Quy định Shopee Affiliate:
- Chỉ tải lên file Excel (.xlsx / .xls)
- Tối đa 10,000 link mỗi file

Cách dùng:
  python3 scripts/export-batch-custom-links.py
  python3 scripts/export-batch-custom-links.py --max-per-file=10000
  python3 scripts/export-batch-custom-links.py --split-even
"""

import sys
import os
import csv
import shutil
import math

try:
    import openpyxl
except ImportError:
    print("❌ Cần cài đặt openpyxl: pip3 install openpyxl")
    sys.exit(1)

def main():
    args = sys.argv[1:]
    input_file = "data/crawler_urls_pending_affiliate.csv"
    template_file = "data/Batch Custom Links.xlsx"
    max_per_file = 10000
    split_even = False
    pending_only = False
    start_part = 1

    for a in args:
        if a.startswith("--input="):
            input_file = a.split("=", 1)[1].strip()
        elif a.startswith("--max-per-file="):
            max_per_file = int(a.split("=", 1)[1].strip())
        elif a.startswith("--start-part="):
            start_part = int(a.split("=", 1)[1].strip())
        elif a == "--pending-only":
            pending_only = True
        elif a == "--split-even":
            split_even = True

    if not os.path.exists(input_file):
        print(f"❌ Không tìm thấy file đầu vào: {input_file}")
        sys.exit(1)

    print(f"📂 Đang đọc dữ liệu từ: {input_file}...")
    with open(input_file, "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        header = next(reader)
        # Cột thứ 4 (index 3) là Link ShopeeFood Gốc, cột thứ 5 (index 4) là Link Affiliate
        urls = []
        for r in reader:
            if len(r) > 3 and r[3].strip():
                if pending_only:
                    aff = r[4].strip() if len(r) > 4 else ""
                    if not aff or ("shope.ee" not in aff and "s.shopee.vn" not in aff):
                        urls.append(r[3].strip())
                else:
                    urls.append(r[3].strip())

    total_urls = len(urls)
    filter_msg = " (chỉ các quán CHƯA CÓ link affiliate)" if pending_only else ""
    print(f"✅ Đã nạp {total_urls:,} links từ CSV{filter_msg}.")

    if total_urls == 0:
        print("⚠️ Không có link nào để xuất.")
        return

    # Xác định cách chia file
    if split_even:
        num_files = math.ceil(total_urls / max_per_file)
        chunk_size = math.ceil(total_urls / num_files)
        chunks = [urls[i:i + chunk_size] for i in range(0, total_urls, chunk_size)]
    else:
        chunks = [urls[i:i + max_per_file] for i in range(0, total_urls, max_per_file)]

    print(f"📦 Sẽ chia thành {len(chunks)} file Excel (tối đa {max_per_file:,} link/file):")

    for chunk_idx, chunk in enumerate(chunks):
        part_num = start_part + chunk_idx
        out_name = f"Batch Custom Links_Part{part_num}.xlsx"
        out_path = os.path.join("data", out_name)

        if os.path.exists(template_file):
            wb = openpyxl.load_workbook(template_file)
            ws = wb.active
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            ws.append(["Liên kết gốc", "Sub_id1", "Sub_id2", "Sub_id3", "Sub_id4", "Sub_id5"])

        for u in chunk:
            ws.append([u, None, None, None, None, None])

        wb.save(out_path)
        print(f"  📄 File Part {part_num}: {out_name} -> {len(chunk):,} link (tổng {ws.max_row:,} dòng)")

        # Tạo các bản sao tên thay thế để người dùng thuận tiện tìm kiếm
        alt_name_space = f"Batch Custom Links {part_num}.xlsx"
        shutil.copyfile(out_path, os.path.join("data", alt_name_space))

        alt_name_singular = f"Batch Custom Link_Part{part_num}.xlsx"
        shutil.copyfile(out_path, os.path.join("data", alt_name_singular))

    print("\n🎉 Hoàn thành xuất các file Excel Batch Custom Links!")
    print("👉 Hãy đăng nhập Shopee Affiliate Portal (https://affiliate.shopee.vn) > Custom Link > Batch")
    print("👉 Tải từng file lên để Shopee chuyển đổi hàng loạt sang link rút gọn (shope.ee / s.shopee.vn).")

if __name__ == "__main__":
    main()
