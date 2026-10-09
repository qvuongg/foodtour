#!/usr/bin/env python3
"""
============================================================================
🏪 EXPORT BATCH CUSTOM LINKS CHO 10 CHUỖI THƯƠNG HIỆU LỚN (SHOPEE AFFILIATE)
============================================================================
Công dụng:
  Đọc toàn bộ 1,530 chi nhánh từ 'data/brand_branches_master.json'
  và xuất ra các file chuẩn mẫu Shopee Affiliate để người dùng upload lên
  https://affiliate.shopee.vn lấy link rút gọn (shope.ee / s.shopee.vn).

Các file được tạo ra trong data/:
  1. data/Batch Custom Links_Brand_Branches_NoSubId.xlsx
     -> Chuẩn 100% mẫu Shopee, để trống Sub_id để tỷ lệ thành công tuyệt đối.
  2. data/Batch Custom Links_Brand_Branches.xlsx
     -> Kèm Sub_id1 (mã thương hiệu) và Sub_id2 (thành phố) để theo dõi hoa hồng.
  3. data/brand_branches_pending_affiliate.csv
     -> File CSV tổng hợp đầy đủ thông tin quán để dễ tra cứu, đối soát.

Cách chạy:
  python3 scripts/export-brand-branches-custom-links.py
============================================================================
"""

import os
import sys
import json
import csv
import shutil

try:
    import openpyxl
except ImportError:
    print("❌ Cần cài đặt openpyxl: pip3 install openpyxl")
    sys.exit(1)

INPUT_JSON = "data/brand_branches_master.json"
TEMPLATE_EXCEL = "data/Batch Custom Links.xlsx"

OUT_EXCEL_NO_SUBID = "data/Batch Custom Links_Brand_Branches_NoSubId.xlsx"
OUT_EXCEL_WITH_SUBID = "data/Batch Custom Links_Brand_Branches.xlsx"
OUT_CSV = "data/brand_branches_pending_affiliate.csv"

# Mã viết tắt sạch cho Sub_id (chỉ chữ và số, không ký tự đặc biệt)
BRAND_SUB_IDS = {
    "highlands": "hl",
    "phuclong": "pl",
    "phela": "phela",
    "katinat": "kat",
    "starbucks": "sbx",
    "mixue": "mixue",
    "tocotoco": "toco",
    "congcaphe": "cong",
    "gongcha": "gc",
    "koithe": "koi",
}

CITY_SUB_IDS = {
    "ho-chi-minh": "hcm",
    "ha-noi": "hn",
    "da-nang": "dn",
}

def main():
    print("================================================================")
    print("📤 XUẤT FILE BATCH CUSTOM LINKS CHO 10 THƯƠNG HIỆU LỚN")
    print("================================================================")

    if not os.path.exists(INPUT_JSON):
        print(f"❌ Không tìm thấy file dữ liệu: {INPUT_JSON}")
        sys.exit(1)

    with open(INPUT_JSON, "r", encoding="utf-8") as f:
        branches = json.load(f)

    total_count = len(branches)
    print(f"📦 Đã đọc {total_count:,} chi nhánh từ {INPUT_JSON}.")

    # 1. Tạo file Excel không có Sub_id (khuyến nghị số 1 cho Shopee Portal)
    print(f"\n1️⃣  Đang tạo {OUT_EXCEL_NO_SUBID} (No Sub ID)...")
    if os.path.exists(TEMPLATE_EXCEL):
        wb_no_sub = openpyxl.load_workbook(TEMPLATE_EXCEL)
        ws_no_sub = wb_no_sub.active
        if ws_no_sub.max_row > 1:
            ws_no_sub.delete_rows(2, ws_no_sub.max_row - 1)
    else:
        wb_no_sub = openpyxl.Workbook()
        ws_no_sub = wb_no_sub.active
        ws_no_sub.title = "Sheet1"
        ws_no_sub.append(["Liên kết gốc", "Sub_id1", "Sub_id2", "Sub_id3", "Sub_id4", "Sub_id5"])

    for b in branches:
        url = b.get("shopeefood_url", "").strip()
        if url:
            ws_no_sub.append([url, None, None, None, None, None])

    wb_no_sub.save(OUT_EXCEL_NO_SUBID)
    print(f"   ✅ Đã ghi {ws_no_sub.max_row - 1:,} links vào {OUT_EXCEL_NO_SUBID}")

    # 2. Tạo file Excel có Sub_id1 (Brand) và Sub_id2 (City)
    print(f"\n2️⃣  Đang tạo {OUT_EXCEL_WITH_SUBID} (With Sub IDs)...")
    if os.path.exists(TEMPLATE_EXCEL):
        wb_sub = openpyxl.load_workbook(TEMPLATE_EXCEL)
        ws_sub = wb_sub.active
        if ws_sub.max_row > 1:
            ws_sub.delete_rows(2, ws_sub.max_row - 1)
    else:
        wb_sub = openpyxl.Workbook()
        ws_sub = wb_sub.active
        ws_sub.title = "Sheet1"
        ws_sub.append(["Liên kết gốc", "Sub_id1", "Sub_id2", "Sub_id3", "Sub_id4", "Sub_id5"])

    for b in branches:
        url = b.get("shopeefood_url", "").strip()
        if url:
            b_id = b.get("brand_id", "")
            c_slug = b.get("city", "")
            sub1 = BRAND_SUB_IDS.get(b_id, b_id[:8])
            sub2 = CITY_SUB_IDS.get(c_slug, c_slug[:4])
            ws_sub.append([url, sub1, sub2, None, None, None])

    wb_sub.save(OUT_EXCEL_WITH_SUBID)
    print(f"   ✅ Đã ghi {ws_sub.max_row - 1:,} links vào {OUT_EXCEL_WITH_SUBID}")

    # 3. Tạo file CSV tổng hợp đối soát (Tên Quán, Thương Hiệu, Thành Phố, Quận, Link Gốc, Link Affiliate)
    print(f"\n3️⃣  Đang tạo {OUT_CSV} (File CSV đối soát)...")
    with open(OUT_CSV, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow([
            "Tên Quán",
            "Thương Hiệu",
            "Thành Phố",
            "Quận",
            "Link ShopeeFood Gốc",
            "Link Affiliate shope.ee (Điền vào đây)"
        ])
        for b in branches:
            writer.writerow([
                b.get("name", ""),
                b.get("brand_name", ""),
                b.get("city", ""),
                b.get("district", ""),
                b.get("shopeefood_url", ""),
                b.get("affiliate_url", "") or ""
            ])
    print(f"   ✅ Đã ghi {total_count:,} dòng vào {OUT_CSV}")

    # Thống kê theo thương hiệu
    print("\n📊 THỐNG KÊ CHI NHÁNH THEO THƯƠNG HIỆU ĐÃ XUẤT:")
    brand_counts = {}
    for b in branches:
        b_name = b.get("brand_name", b.get("brand_id", "Khác"))
        brand_counts[b_name] = brand_counts.get(b_name, 0) + 1

    for name, cnt in sorted(brand_counts.items(), key=lambda x: -x[1]):
        print(f"   • {name.ljust(20)}: {cnt:,} quán")

    print("\n================================================================")
    print("🎉 HOÀN THÀNH XUẤT FILE GỬI LÊN SHOPEE AFFILIATE!")
    print("================================================================")
    print("👉 BƯỚC TIẾP THEO DÀNH CHO BẠN:")
    print("  1. Truy cập Shopee Affiliate Portal: https://affiliate.shopee.vn")
    print("  2. Vào mục: Link tùy chỉnh (Custom Link) > Tạo link hàng loạt (Batch Link Generation)")
    print(f"  3. Tải lên file: {OUT_EXCEL_NO_SUBID} (hoặc {OUT_EXCEL_WITH_SUBID})")
    print("  4. Đợi Shopee xử lý xong (thường mất 1 - 2 phút) và tải file kết quả về máy.")
    print("     (File kết quả thường có tên dạng: AffiliateBatchCustomLinks_YYYYMMDD_HHMM.csv / .xlsx)")
    print("  5. Bỏ file kết quả đó vào thư mục data/ của dự án để hệ thống tự động gộp link affiliate vào database!")
    print("================================================================")

if __name__ == "__main__":
    main()
